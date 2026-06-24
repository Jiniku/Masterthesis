#!/usr/bin/env python3
"""Build quizgen_slides.pptx — a 'how it works' deck for the quizgen project.

quizgen is a selection-only tool: it assembles blueprint-compliant exams from
an existing question bank (no generation, no LLM, no PDF ingestion).

Run:  python build_slides.py   (writes docs/quizgen_slides.pptx)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor

# ── palette ───────────────────────────────────────────────────────────
NAVY   = RGBColor(0x1F, 0x2D, 0x3D)
BLUE   = RGBColor(0x31, 0x82, 0xBD)
BLUEL  = RGBColor(0xDE, 0xEB, 0xF7)
GREEN  = RGBColor(0x31, 0xA3, 0x54)
GREENL = RGBColor(0xE5, 0xF5, 0xE0)
RED    = RGBColor(0xDE, 0x2D, 0x26)
REDL   = RGBColor(0xFD, 0xE0, 0xDD)
GREY   = RGBColor(0x63, 0x63, 0x63)
GREYL  = RGBColor(0xF0, 0xF0, 0xF2)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x33, 0x33, 0x33)

FONT = "Calibri"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ── helpers ───────────────────────────────────────────────────────────
def slide():
    return prs.slides.add_slide(BLANK)


def _set(frame, size, color, bold=False, align=PP_ALIGN.LEFT, font=FONT, italic=False):
    for p in frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.name = font
            r.font.color.rgb = color


def textbox(s, x, y, w, h, text, size=18, color=DARK, bold=False,
            align=PP_ALIGN.LEFT, font=FONT, italic=False, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.text = text
    _set(tf, size, color, bold, align, font, italic)
    return tb


def box(s, x, y, w, h, text, fill, line, txtcolor=DARK, size=14, bold=True,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, font=FONT, line_w=1.25):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03); tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = text
    _set(tf, size, txtcolor, bold, PP_ALIGN.CENTER, font)
    return sp


def arrow(s, x1, y1, x2, y2, color=GREY, w=2.0, dash=False):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.line.color.rgb = color; cn.line.width = Pt(w)
    le = cn.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = le.makeelement(qn('a:tailEnd'),
                          {'type': 'triangle', 'w': 'med', 'len': 'med'})
    le.append(tail)
    if dash:
        d = le.makeelement(qn('a:prstDash'), {'val': 'dash'})
        le.insert(0, d)
    return cn


def header(s, title, kicker=None):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.15))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    bar.shadow.inherit = False
    textbox(s, Inches(0.55), Inches(0.18), Inches(12.2), Inches(0.8),
            title, size=30, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        textbox(s, Inches(0.58), Inches(0.02), Inches(12), Inches(0.3),
                kicker, size=12, color=BLUEL, bold=True)


def bullets(s, x, y, w, h, items, size=18, gap=6):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        lvl = 0
        if isinstance(it, tuple):
            it, lvl = it
        p.level = lvl
        run = p.add_run(); run.text = ("• " if lvl == 0 else "– ") + it
        run.font.size = Pt(size - lvl * 2)
        run.font.name = FONT
        run.font.color.rgb = DARK if lvl == 0 else GREY
        p.space_after = Pt(gap)
    return tb


def chip(s, x, y, w, text, fill, line, txt=DARK, size=11):
    return box(s, x, y, w, Inches(0.42), text, fill, line, txt, size=size,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE)


# ══════════════════════════════════════════════════════════════════════
# 1 — TITLE
# ══════════════════════════════════════════════════════════════════════
s = slide()
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
bg.shadow.inherit = False
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.7), SW, Inches(0.12))
band.fill.solid(); band.fill.fore_color.rgb = BLUE; band.line.fill.background()
band.shadow.inherit = False
textbox(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.2),
        "quizgen", size=66, color=WHITE, bold=True)
textbox(s, Inches(0.85), Inches(3.25), Inches(11.7), Inches(1.3),
        "Automated Test Assembly from a question bank\n"
        "Select blueprint-compliant exams — no generation, no LLM",
        size=26, color=BLUEL)
textbox(s, Inches(0.85), Inches(5.0), Inches(11.7), Inches(1.4),
        "Question bank (JSON)   •   Blueprint (YAML + flags)   •   "
        "Greedy / MIP selection\nMaster's Thesis — KWARC Group, FAU "
        "Erlangen-Nürnberg", size=16, color=WHITE)

# ══════════════════════════════════════════════════════════════════════
# 2 — THE IDEA
# ══════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Selection as exact optimisation", "WHAT THE TOOL DOES")
textbox(s, Inches(0.55), Inches(1.35), Inches(12.2), Inches(0.8),
        "You bring a bank of tagged questions. quizgen picks the exam that "
        "matches the professor's blueprint — exactly.", size=18, italic=True, color=GREY)

box(s, Inches(0.7), Inches(2.35), Inches(5.3), Inches(1.0),
    "INPUT  ·  question bank (JSON) with metadata + a blueprint",
    GREENL, GREEN, DARK, size=15)
box(s, Inches(7.3), Inches(2.35), Inches(5.3), Inches(1.0),
    "OUTPUT  ·  one or more disjoint, blueprint-compliant exam versions",
    BLUEL, BLUE, DARK, size=15)
arrow(s, Inches(6.0), Inches(2.85), Inches(7.3), Inches(2.85), GREY, 2.5)

bullets(s, Inches(0.7), Inches(3.75), Inches(12), Inches(3.1), [
    "Hitting counts exactly (per chapter, per difficulty, N versions) is a "
    "combinatorial optimisation problem — not a prompting problem.",
    "Solved with a 0–1 Mixed-Integer Program: exact, auditable, optimal under a quality objective.",
    "Parallel versions are produced in one solve and provably share no questions.",
    "Every exam embeds the resolved blueprint → self-describing & reproducible.",
    "No LLM, no PDF ingestion, no network — runs fully offline.",
], size=18, gap=10)

# ══════════════════════════════════════════════════════════════════════
# 3 — ARCHITECTURE / PIPELINE DIAGRAM
# ══════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Architecture: the pipeline at a glance", "POOL → SELECT → VERIFY")

bw, bh = Inches(2.2), Inches(1.2)
y0 = Inches(2.9)
xs = [Inches(0.7), Inches(3.5), Inches(6.5), Inches(9.5)]
labels = [("question bank\n(JSON)", GREENL, GREEN),
          ("Deduplicate\n(optional)", BLUEL, BLUE),
          ("Assemble\n(greedy / MIP)", BLUEL, BLUE),
          ("Validate\n(optional)", BLUEL, BLUE)]
for i, (lab, f, l) in enumerate(labels):
    shape = MSO_SHAPE.RECTANGLE if i == 0 else MSO_SHAPE.ROUNDED_RECTANGLE
    box(s, xs[i], y0, bw, bh, lab, f, l, DARK, size=14, shape=shape)
    if i < len(labels) - 1:
        arrow(s, xs[i] + bw, y0 + bh // 2, xs[i + 1], y0 + bh // 2, GREY, 1.8)
# output
out_x = Inches(11.9)
box(s, out_x, y0, Inches(1.2), bh, "exam\nversions", GREENL, GREEN, DARK, size=11,
    shape=MSO_SHAPE.RECTANGLE)
arrow(s, xs[3] + bw, y0 + bh // 2, out_x, y0 + bh // 2, GREEN, 1.6)
# blueprint feeds assemble
box(s, xs[2], Inches(1.5), bw, Inches(0.7), "blueprint\n(YAML + flags)",
    REDL, RED, DARK, size=12, shape=MSO_SHAPE.RECTANGLE)
arrow(s, xs[2] + bw // 2, Inches(2.2), xs[2] + bw // 2, y0, RED, 1.6)

# legend
chip(s, Inches(0.7), Inches(5.0), Inches(2.0), "data", GREENL, GREEN)
chip(s, Inches(2.9), Inches(5.0), Inches(2.2), "processing", BLUEL, BLUE)
chip(s, Inches(5.3), Inches(5.0), Inches(2.6), "professor input", REDL, RED)
textbox(s, Inches(0.7), Inches(5.7), Inches(12), Inches(1.2),
        "Dedup and validation are opt-in (--dedup / --validate). The blueprint "
        "drives assembly; each stage is one module in the quizgen package.",
        size=15, color=GREY, italic=True)

# ══════════════════════════════════════════════════════════════════════
# 4 — DATA MODEL / INPUT
# ══════════════════════════════════════════════════════════════════════
s = slide()
header(s, "The input: a tagged question bank", "schema.py — Pydantic v2")
bullets(s, Inches(0.7), Inches(1.45), Inches(5.7), Inches(5), [
    "A JSON Quiz object (or a bare list) of Question items.",
    "Pydantic v2 models are the single source of truth; JSON Schema is exported, not hand-written.",
    "Metadata is exactly what selection reasons over:",
    ("chapter → per-chapter quotas", 1),
    ("difficulty → easy/medium/hard mix", 1),
    ("qtype → allowed-type filter", 1),
    ("bloom_level → quality objective", 1),
    ("stem/answer/options → duplicate detection", 1),
], size=17, gap=7)

code = (
    "{\n"
    '  "chapter": 1,\n'
    '  "section": "1.2",\n'
    '  "qtype": "mcq",\n'
    '  "bloom_level": "understand",\n'
    '  "difficulty": "medium",\n'
    '  "stem": "Explain how A* uses ...",\n'
    '  "options": ["A) ...", "B) ...",\n'
    '              "C) ...", "D) ..."],\n'
    '  "answer": "A",\n'
    '  "explanation": "A* is optimal ...",\n'
    '  "source_ref": "ch1_sec2_q1"\n'
    "}"
)
cb = box(s, Inches(6.7), Inches(1.45), Inches(6.0), Inches(4.4), code,
         GREYL, GREY, DARK, size=14, bold=False, shape=MSO_SHAPE.RECTANGLE, font=MONO)
for p in cb.text_frame.paragraphs:
    p.alignment = PP_ALIGN.LEFT

# ══════════════════════════════════════════════════════════════════════
# 5 — BLUEPRINT
# ══════════════════════════════════════════════════════════════════════
s = slide()
header(s, "The blueprint: a table of specifications", "blueprint.py")
bp = (
    "total_questions: 20\n"
    "chapters: {1: 50%, 2: 30%, 3: 20%}\n"
    "difficulty_mix: {easy:.4, medium:.4, hard:.2}\n"
    "allowed_qtypes: [mcq, short_answer, ...]\n"
    "versions: 2\n"
    "selector: mip\n"
    "dedup_similarity: 0.85"
)
box(s, Inches(0.7), Inches(1.5), Inches(5.9), Inches(2.6), bp,
    REDL, RED, DARK, size=14, bold=False, font=MONO, shape=MSO_SHAPE.RECTANGLE)
bullets(s, Inches(0.7), Inches(4.4), Inches(5.9), Inches(2.6), [
    "Counts may be ints, fractions, or \"NN%\".",
    "resolve() → exact integer counts that sum to the total (largest-remainder).",
], size=15, gap=8)

textbox(s, Inches(7.0), Inches(1.45), Inches(5.7), Inches(0.5),
        "YAML and/or CLI flags — flags win", size=17, bold=True, color=BLUE)
bullets(s, Inches(7.0), Inches(2.0), Inches(5.7), Inches(2.0), [
    "Keep a base blueprint.yaml, override per run.",
    "build_blueprint loads YAML → overwrites field-by-field with provided flags → validates.",
], size=15, gap=8)
box(s, Inches(7.0), Inches(4.0), Inches(5.8), Inches(1.6),
    "python -m quizgen --pool questions.json \\\n"
    "  --blueprint blueprint.yaml \\\n"
    "  --versions 3 --selector greedy",
    NAVY, NAVY, WHITE, size=13, bold=False, font=MONO, shape=MSO_SHAPE.RECTANGLE)

# ══════════════════════════════════════════════════════════════════════
# 6 — ATA: GREEDY vs MIP + FORMULATION
# ══════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Automated Test Assembly: greedy vs. MIP", "assemble.py")
box(s, Inches(0.7), Inches(1.45), Inches(5.7), Inches(0.55),
    "GreedySelector — fast baseline", BLUE, BLUE, WHITE, size=14)
bullets(s, Inches(0.75), Inches(2.1), Inches(5.6), Inches(2.0), [
    "Per-chapter fill, two passes (difficulty quota, then remainder).",
    "Skips duplicates & already-used items.",
    "Fast; no optimality guarantee.",
], size=14, gap=6)
box(s, Inches(0.7), Inches(4.1), Inches(5.7), Inches(0.55),
    "MIPSelector — exact optimisation", BLUE, BLUE, WHITE, size=14)
bullets(s, Inches(0.75), Inches(4.75), Inches(5.6), Inches(2.2), [
    "0–1 program via PuLP / CBC.",
    "Exact per-chapter AND difficulty counts.",
    "Disjoint versions; forbids dup pairs.",
    "Infeasible → relax difficulty (logged), never chapter counts.",
], size=14, gap=6)

mip = (
    "max  Σ q_i · x_iv          (quality)\n"
    "Σ_{chap(i)=c} x_iv = n_c    (per chapter)\n"
    "Σ_{diff(i)=d} x_iv = m_d    (difficulty)\n"
    "Σ_v  x_iv ≤ 1              (disjoint)\n"
    "Σ_v x_iv + Σ_v x_jv ≤ 1    (no dup pair)\n"
    "x_iv ∈ {0, 1}"
)
box(s, Inches(6.8), Inches(1.7), Inches(6.0), Inches(3.0), mip,
    GREYL, GREY, DARK, size=15, bold=False, font=MONO, shape=MSO_SHAPE.RECTANGLE)
textbox(s, Inches(6.8), Inches(4.9), Inches(6.0), Inches(1.6),
        "One solve produces all parallel versions at once — provably disjoint "
        "and each individually blueprint-compliant.", size=15, color=GREY, italic=True)

# ══════════════════════════════════════════════════════════════════════
# 7 — DEDUP + VALIDATION
# ══════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Deduplication & validation", "similarity.py • dedup.py • validate.py")
textbox(s, Inches(0.7), Inches(1.3), Inches(6), Inches(0.4),
        "DEDUP (--dedup)", size=17, color=BLUE, bold=True)
bullets(s, Inches(0.7), Inches(1.75), Inches(5.8), Inches(2.4), [
    "Token-Jaccard over stem+answer(+options) — offline, always on.",
    "Optional embedding cosine — catches paraphrases.",
    "Drops dups from the pool; MIP also forbids dup pairs across versions.",
], size=15, gap=7)
box(s, Inches(0.7), Inches(4.4), Inches(5.7), Inches(1.1),
    "sim(q₁,q₂) = |T(q₁) ∩ T(q₂)| / |T(q₁) ∪ T(q₂)|",
    GREYL, GREY, DARK, size=14, bold=False, font=MONO, shape=MSO_SHAPE.RECTANGLE)

textbox(s, Inches(7.0), Inches(1.3), Inches(6), Inches(0.4),
        "VALIDATE (--validate)", size=17, color=GREEN, bold=True)
quad = [("Schema validity", GREENL, GREEN), ("Blueprint compliance", BLUEL, BLUE),
        ("Grounding coverage", GREENL, GREEN), ("Duplicate rate", BLUEL, BLUE)]
for i, (t, f, l) in enumerate(quad):
    x = Inches(7.0) + (i % 2) * Inches(2.95)
    y = Inches(1.85) + (i // 2) * Inches(1.0)
    box(s, x, y, Inches(2.8), Inches(0.85), t, f, l, DARK, size=13)
textbox(s, Inches(7.0), Inches(4.0), Inches(5.8), Inches(1.6),
        "Overall PASS only if no schema-invalid items and (if checked) the "
        "blueprint matches exactly. Exit code feeds CI.", size=14, color=GREY, italic=True)

# ══════════════════════════════════════════════════════════════════════
# 8 — CLI / USAGE
# ══════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Using it", "python -m quizgen")
textbox(s, Inches(0.7), Inches(1.35), Inches(12), Inches(0.4),
        "Pure command line (no YAML):", size=16, bold=True, color=DARK)
box(s, Inches(0.7), Inches(1.8), Inches(12.0), Inches(1.25),
    "python -m quizgen --pool questions.json \\\n"
    "    --total 20 --chapters 1:10,2:6,3:4 \\\n"
    "    --difficulty 0.4,0.4,0.2 --versions 2 --selector mip",
    NAVY, NAVY, WHITE, size=15, bold=False, font=MONO, shape=MSO_SHAPE.RECTANGLE)
textbox(s, Inches(0.7), Inches(3.3), Inches(12), Inches(0.4),
        "Blueprint file + dedup + per-version validation report:", size=16, bold=True, color=DARK)
box(s, Inches(0.7), Inches(3.75), Inches(12.0), Inches(1.0),
    "python -m quizgen --pool questions.json --blueprint blueprint.yaml \\\n"
    "    --dedup --validate --out-dir out",
    NAVY, NAVY, WHITE, size=15, bold=False, font=MONO, shape=MSO_SHAPE.RECTANGLE)
bullets(s, Inches(0.7), Inches(5.1), Inches(12), Inches(1.8), [
    "Writes one Quiz JSON per version (each embeds the resolved blueprint).",
    "Standalone entry points too: python -m quizgen.assemble, python -m quizgen.validate.",
    "Exits non-zero on any non-compliant version → composes in scripts & CI.",
], size=16, gap=8)

# ══════════════════════════════════════════════════════════════════════
# 9 — SOURCE ATTRIBUTION
# ══════════════════════════════════════════════════════════════════════
s = slide()
header(s, "Where each technique comes from", "TECHNIQUE → SOURCE LINEAGE")
rows = [
    ("Test assembly as 0–1 MIP", "van der Linden 2005 — optimal test design"),
    ("Quality-weighted selection", "van der Linden 2005 (objective in ATA)"),
    ("Test blueprint", "classical table-of-specifications design"),
    ("Bloom-level metadata", "Anderson & Krathwohl 2001 (revised Bloom)"),
    ("Semantic dedup (embeddings)", "Sentence-BERT (Reimers & Gurevych 2019)"),
    ("Largest-remainder counts", "Hamilton apportionment (engineering)"),
    ("Greedy baseline, token Jaccard", "project-specific engineering"),
]
ty = Inches(1.5)
rh = Inches(0.58)
box(s, Inches(0.7), ty, Inches(5.4), rh, "Implemented in code", BLUE, BLUE, WHITE, size=14)
box(s, Inches(6.2), ty, Inches(6.5), rh, "Source lineage", BLUE, BLUE, WHITE, size=14)
for i, (a, b) in enumerate(rows):
    y = ty + rh * (i + 1)
    f = WHITE if i % 2 else GREYL
    box(s, Inches(0.7), y, Inches(5.4), rh, a, f, GREY, DARK, size=12, bold=False,
        shape=MSO_SHAPE.RECTANGLE)
    box(s, Inches(6.2), y, Inches(6.5), rh, b, f, GREY, DARK, size=12, bold=False,
        shape=MSO_SHAPE.RECTANGLE)
textbox(s, Inches(0.7), Inches(6.95), Inches(12.4), Inches(0.4),
        "Note: replace these keys with the thesis's exact bibliography.",
        size=11, color=RED, italic=True)

# ══════════════════════════════════════════════════════════════════════
# 10 — TAKEAWAYS
# ══════════════════════════════════════════════════════════════════════
s = slide()
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
bg.shadow.inherit = False
textbox(s, Inches(0.8), Inches(0.7), Inches(12), Inches(0.9),
        "Takeaways", size=40, color=WHITE, bold=True)
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.65),
                          Inches(3.2), Inches(0.08))
band.fill.solid(); band.fill.fore_color.rgb = BLUE; band.line.fill.background()
band.shadow.inherit = False
items = [
    "Scope: SELECT an exam from a question bank — no generation, no LLM.",
    "Blueprint = exact spec; counts resolve via largest-remainder.",
    "Constraints from YAML and/or CLI flags (flags override).",
    "A 0–1 MIP gives exact, optimal, disjoint multi-version assembly.",
    "Optional dedup + layered validation (schema, blueprint, dups).",
    "Runs fully offline; greedy baseline behind the same interface.",
]
tb = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.6), Inches(4.8))
tf = tb.text_frame; tf.word_wrap = True
for i, it in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r = p.add_run(); r.text = "→  " + it
    r.font.size = Pt(22); r.font.name = FONT; r.font.color.rgb = WHITE
    p.space_after = Pt(16)

out = "docs/quizgen_slides.pptx"
prs.save(out)
print("saved", out, "with", len(prs.slides._sldIdLst), "slides")
